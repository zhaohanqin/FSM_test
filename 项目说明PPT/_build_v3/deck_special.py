# -*- coding: utf-8 -*-
"""deck_special.py — 封面 / 目录 / 章节页 / 结束页 等特殊版式。"""

from deck_kit import *
from deck_kit import _no_shadow, _set_ea, _hl_line
from pptx.enum.shapes import MSO_SHAPE


def _wave(s, y_center, amp, color="1B4A73", w_pt=0.9, op=60, phase=0.0,
          x0=0.0, x1=SW, n=220):
    """在深色底上画一条淡色正弦波，增加科技感。"""
    import math
    pts = []
    for i in range(n + 1):
        x = x0 + (x1 - x0) * i / n
        y = y_center + amp * math.sin(2 * math.pi * (i / n) * 2.4 + phase)
        pts.append((x, y))
    ff = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
    ff.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    shp = ff.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = RGBColor.from_string(color)
    shp.line.width = Pt(w_pt)
    shp.shadow.inherit = False
    _no_shadow(shp)
    return shp


def cover(s, title_lines, subtitle, eyebrow_txt, formula=None, chips_items=None,
          note=None, section=""):
    bg_dark(s, deep=True)
    left_bar(s, "cyan_br")
    # 背景波纹（放在副标题与底部 chip 之间的空白带）
    _wave(s, 4.30, 0.26, "1B4A73", 0.9, phase=0.0)
    _wave(s, 4.72, 0.17, "173F63", 0.8, phase=1.1)
    _wave(s, 5.14, 0.26, "15395A", 0.7, phase=2.2)

    text(s, ML, 1.12, CW, 0.26, eyebrow_txt, size=11, bold=True, color="cyan_br")

    y = 1.62
    for i, (txt, col) in enumerate(title_lines):
        text(s, ML, y, 8.4, 0.85, txt, size=40, bold=True, color=col, spacing=0.95)
        y += 0.92
    text(s, ML, y + 0.06, 8.3, 0.4, subtitle, size=13, color="8AA4B8", spacing=1.1)

    if formula:
        fx = 8.95
        rrect(s, fx, 1.78, 3.78, 1.42, fill="0D2138", line="1D5C7A", radius=0.07,
              line_w=1.1)
        text(s, fx + 0.24, 1.78, 3.3, 1.42, formula, size=19, bold=True,
             color="22D3EE", anchor="m", align="l", spacing=1.05)
        if note:
            text(s, fx + 0.02, 3.34, 3.74, 0.3, note, size=10, color="8AA4B8",
                 align="r")

    if chips_items:
        chips(s, ML, 5.55, chips_items, w=CW - 0.02, h=0.5, gap=0.36,
              fill="0D2138", color="8FB6CC", size=11)
    return s


def chapter(s, num_cn, title, en, blurb, section, pageno, total=None):
    bg_dark(s)
    _wave(s, 5.72, 0.24, "1B4A73", 0.9, phase=0.4)
    _wave(s, 6.24, 0.16, "173F63", 0.8, phase=1.6)
    # 左侧装饰块
    rect(s, 1.03, 2.86, 1.62, 0.34, fill="1B4A73")
    rect(s, 1.03, 4.03, 1.62, 0.34, fill="1B4A73")
    # CHAPTER 字标
    t = text(s, ML, 2.28, 4.0, 0.3, "C H A P T E R", size=11.5, bold=True,
             color="22D3EE")
    for r in t.text_frame.paragraphs[0].runs:
        r._r.get_or_add_rPr().set("spc", "300")
    text(s, 2.95, 2.62, 9.6, 0.95, title, size=40, bold=True, color="white",
         spacing=0.95)
    text(s, 2.95, 3.66, 9.6, 0.34, en, size=14, color="22D3EE")
    line_h(s, 2.95, 4.24, 8.4, "2A5A80", 1.0)
    text(s, 2.95, 4.46, 9.4, 0.66, blurb, size=13, color="D7E6F2", spacing=1.3)
    footer(s, section, pageno, total)


def toc(s, items, flow, flow_label, section, pageno, total=None):
    """目录页：上方主线流程条 + 双列章节卡。"""
    bg_light(s)
    left_bar(s, "teal")
    header(s, "讲解路线 · ROADMAP", "本次讲解的主线与章节安排",
           "一条主线贯穿：从条纹图像到三维点云；八章依次展开。")
    text(s, ML, 1.94, CW, 0.24, flow_label, size=11.5, bold=True, color="teal")
    chips(s, ML, 2.22, flow, w=CW, h=0.5, gap=0.30, fill="fill_2",
          color="navy", size=11.5, arrow_color="cyan")
    y0 = 3.00
    n = len(items)
    rows = (n + 1) // 2
    cw = (CW - 0.34) / 2
    ch = (BODY_BOTTOM - y0 - 0.14 * (rows - 1)) / rows
    nums = ["一", "二", "三", "四", "五", "六", "七", "八"]
    accents = ["teal", "amber", "teal", "amber", "teal", "amber", "teal", "amber"]
    for i, (title, desc) in enumerate(items):
        c, r = divmod(i, rows)
        cx = ML + c * (cw + 0.34)
        cy = y0 + r * (ch + 0.14)
        rrect(s, cx, cy, cw, ch, fill="white", line="line", radius=0.07)
        rect(s, cx, cy + 0.07, 0.06, ch - 0.14, fill=accents[i % 2])
        rect(s, cx + 0.20, cy + ch / 2 - 0.17, 0.34, 0.34, fill="navy",
             shape=MSO_SHAPE.OVAL)
        text(s, cx + 0.20, cy + ch / 2 - 0.12, 0.34, 0.26, nums[i], size=11.5,
             bold=True, color="white", align="c")
        text(s, cx + 0.66, cy + 0.12, cw - 0.86, 0.3, title, size=13.5,
             bold=True, color="navy")
        text(s, cx + 0.66, cy + 0.42, cw - 0.86, ch - 0.54, desc, size=10.5,
             color="gray", spacing=1.15)
    footer(s, section, pageno, total)


def ending(s, title, subtitle, points, section, pageno, total=None):
    bg_dark(s, deep=True)
    text(s, ML, 1.30, CW, 0.28, "SUMMARY", size=11, bold=True, color="cyan_br")
    text(s, ML, 1.70, 9.6, 0.8, title, size=34, bold=True, color="white")
    text(s, ML, 2.60, 9.8, 0.34, subtitle, size=13, color="8AA4B8")
    line_h(s, ML, 3.12, 11.2, "2A5A80", 1.0)
    y = 3.36
    for head, desc in points:
        rect(s, ML, y + 0.06, 0.055, 0.3, fill="cyan")
        text(s, ML + 0.22, y, 2.6, 0.3, head, size=12.5, bold=True, color="22D3EE")
        text(s, ML + 2.95, y + 0.015, 8.6, 0.3, desc, size=11.5, color="D7E6F2")
        y += 0.44
    # 波纹放在条目下方的空白带
    wy = max(y + 0.30, 5.60)
    if wy + 0.30 < FOOTER_Y:
        _wave(s, wy, 0.22, "1B4A73", 0.9, phase=0.7)
        _wave(s, wy + 0.42, 0.15, "173F63", 0.8, phase=1.9)
    footer(s, section, pageno, total)
